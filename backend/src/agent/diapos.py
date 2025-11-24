from pptx import Presentation
from pptx.util import Inches, Pt
import os

prs = Presentation()

def add_title_slide(title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    if subtitle is not None:
        slide.placeholders[1].text = subtitle

def add_bullets_slide(title, bullets, font_size=24):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        for run in p.runs:
            run.font.size = Pt(font_size)

# DIAPO 1
add_title_slide("Properties of the Fourier Transform", "Based on the video lecture")

# DIAPO 2
add_bullets_slide(
    "Reminder: Definition of Fourier Transform",
    [
        "Given a signal x(t), its Fourier Transform is:",
        "X(ω) = ∫_{-∞}^{∞} x(t) e^{-jωt} dt",
        "",
        "Inverse Fourier Transform:",
        "x(t) = (1/2π) ∫_{-∞}^{∞} X(ω) e^{+jωt} dω"
    ],
    font_size=22
)

# DIAPO 3
add_bullets_slide(
    "Main Properties We Will Cover",
    [
        "1) Linearity",
        "2) Time Shift (shifted signals)",
        "3) Modulation / Frequency Shift",
        "4) Convolution",
        "5) Derivative (and N-th derivative)",
        "6) Duality"
    ],
    font_size=24
)

# DIAPO 4
add_bullets_slide(
    "Property 1: Linearity",
    [
        "If we have a linear combination of signals:",
        "α₁ x₁(t) + α₂ x₂(t)",
        "",
        "Then the Fourier Transform is the same combination:",
        "α₁ X₁(ω) + α₂ X₂(ω)"
    ],
    font_size=22
)

# DIAPO 5
add_bullets_slide(
    "Proof of Linearity (as in the video)",
    [
        "F{α₁x₁(t) + α₂x₂(t)}",
        "= ∫_{-∞}^{∞} [α₁x₁(t) + α₂x₂(t)] e^{-jωt} dt",
        "",
        "Use linearity of the integral:",
        "= α₁ ∫ x₁(t)e^{-jωt} dt  +  α₂ ∫ x₂(t)e^{-jωt} dt",
        "",
        "Recognize transforms:",
        "= α₁ X₁(ω) + α₂ X₂(ω)"
    ],
    font_size=20
)

# DIAPO 6
add_bullets_slide(
    "Property 2: Time Shift",
    [
        "Consider a shifted signal:",
        "x(t - t₀)",
        "",
        "Fourier Transform becomes:",
        "F{x(t - t₀)} = e^{-jωt₀} X(ω)"
    ],
    font_size=22
)

# DIAPO 7
add_bullets_slide(
    "Proof of Time Shift (as in the video)",
    [
        "F{x(t - t₀)} = ∫ x(t - t₀) e^{-jωt} dt",
        "",
        "Rewrite exponential:",
        "e^{-jωt} = e^{-jω(t - t₀)} e^{-jωt₀}",
        "",
        "Bring constant outside:",
        "= e^{-jωt₀} ∫ x(t - t₀) e^{-jω(t - t₀)} dt",
        "",
        "Change variable τ = t - t₀, dτ = dt:",
        "= e^{-jωt₀} ∫ x(τ) e^{-jωτ} dτ",
        "",
        "Recognize transform:",
        "= e^{-jωt₀} X(ω)"
    ],
    font_size=19
)

# DIAPO 8
add_bullets_slide(
    "Property 3: Modulation / Frequency Shift",
    [
        "If a signal is multiplied by a complex exponential:",
        "e^{jω₀t} x(t)",
        "",
        "Then its spectrum shifts:",
        "F{e^{jω₀t} x(t)} = X(ω - ω₀)"
    ],
    font_size=22
)

# DIAPO 9
add_bullets_slide(
    "Property 4: Convolution",
    [
        "If two signals are convolved in time:",
        "x(t) * y(t)",
        "",
        "Then in frequency we get multiplication:",
        "F{x(t) * y(t)} = X(ω) Y(ω)",
        "",
        "This is one of the most important properties."
    ],
    font_size=22
)

# DIAPO 10
add_bullets_slide(
    "Property 5: Derivative",
    [
        "First derivative:",
        "F{dx(t)/dt} = jω X(ω)",
        "",
        "General N-th derivative:",
        "F{dⁿx(t)/dtⁿ} = (jω)ⁿ X(ω)"
    ],
    font_size=22
)

# DIAPO 11
add_bullets_slide(
    "Property 6: Duality",
    [
        "If we have a transform pair:",
        "x(t) ↔ X(ω)",
        "",
        "Then we get another pair for free:",
        "X(t) ↔ 2π x(-ω)",
        "",
        "Knowing one pair gives you another one."
    ],
    font_size=22
)

# DIAPO 12
add_bullets_slide(
    "Closing Remark",
    [
        "These are the key Fourier Transform properties you need to master.",
        "Their proofs are straightforward applications of integral properties.",
        "Practice proving them to become familiar with the Fourier integral."
    ],
    font_size=24
)

# (Opcional) Si quieres poner tu captura, guarda la imagen como:
# "captura_video.png" en la misma carpeta del script.
img_path = "captura_video.png"
if os.path.exists(img_path):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    slide.shapes.title.text = "Lecture Snapshot (from the video)"
    slide.shapes.add_picture(img_path, Inches(1), Inches(1.5), width=Inches(8))

prs.save("Fourier_Properties_Video.pptx")
print("✅ PPT generado: Fourier_Properties_Video.pptx")
